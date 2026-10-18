"""Word report in the dark slate theme; appends to templates/report_template.docx when present."""

import os
from datetime import datetime

from vuln_reports.reporting.chart import generate_severity_chart
from vuln_reports.reporting.common import TEMPLATES_DIR

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def save_docx_report(vulnerabilities, cve_to_hosts, scanner_type, target_file, output_file):
    if not HAS_DOCX:
        print("[!] python-docx not installed. DOCX report skipped.")
        return

    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    # --- TEMPLATE LOGIC ---
    template_path = os.path.join(TEMPLATES_DIR, "report_template.docx")
    if os.path.exists(template_path):
        doc = Document(template_path)
        print(f"[*] Using custom Word template: {template_path}")
    else:
        doc = Document()

    # Configure page margins for a cleaner dashboard look (0.75 in margins)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # --- GLOBAL DARK THEME ---
    # 1. Set Page Background Color (#0F172A)
    shading_elm_2 = parse_xml(r'<w:background {} w:color="0F172A"/>'.format(nsdecls("w")))
    doc.element.insert(0, shading_elm_2)

    # 2. Enable background color display in Word settings
    settings = doc.settings.element
    display_bg = parse_xml(r"<w:displayBackgroundShape {}/>".format(nsdecls("w")))
    settings.append(display_bg)

    # Helper to set cell shading (background color)
    def set_cell_shading(cell, color_hex):
        shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), color_hex))
        cell._tc.get_or_add_tcPr().append(shading)

    # Helper to set cell borders
    def set_cell_borders(cell, color_hex="334155", sz="4", val="single"):
        borders = parse_xml(
            r'<w:tcBorders {}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>'.format(
                nsdecls("w"), val=val, sz=sz, color=color_hex
            )
        )
        cell._tc.get_or_add_tcPr().append(borders)

    # Helper to format fonts and colors
    def format_run(run, font_name="Segoe UI", size_pt=10, bold=False, italic=False, color_rgb=(248, 250, 252)):
        run.font.name = font_name
        # Apply ASCII/Hansi/CS font settings in oxml to prevent fallback
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(r'<w:rFonts {} w:ascii="{f}" w:hAnsi="{f}" w:cs="{f}"/>'.format(nsdecls("w"), f=font_name))
        rPr.append(rFonts)
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        return run

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    format_run(
        title_p.add_run("Vulnerability Intelligence Report"),
        font_name="Segoe UI",
        size_pt=24,
        bold=True,
        color_rgb=(59, 130, 246),
    )

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    format_run(
        meta_p.add_run(
            f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Scanner: {scanner_type} | Target: {os.path.basename(target_file)}"
        ),
        font_name="Segoe UI",
        size_pt=9.5,
        italic=True,
        color_rgb=(148, 163, 184),
    )

    # Summary Stats Section
    h_exec = doc.add_paragraph()
    h_exec.paragraph_format.space_after = Pt(8)
    format_run(
        h_exec.add_run("Executive Summary"), font_name="Segoe UI", size_pt=16, bold=True, color_rgb=(248, 250, 252)
    )

    stats = {
        "HOSTS": len(set(h[0] for hosts in cve_to_hosts.values() for h in hosts)),
        "Unique CVEs": len(vulnerabilities),
        "Critical Risks": sum(
            1
            for v in vulnerabilities.values()
            if isinstance(v.get("cvss_score"), (int, float)) and v["cvss_score"] >= 9
        ),
        "KEV Exploited": sum(1 for v in vulnerabilities.values() if v.get("cisa_kev") == "Yes"),
    }

    # Generate 4 distinct columns for cards
    stats_table = doc.add_table(rows=1, cols=4)
    stats_table.autofit = False
    col_widths = [Inches(1.75), Inches(1.75), Inches(1.75), Inches(1.75)]

    hdr_cells = stats_table.rows[0].cells
    for i, (label, val) in enumerate(stats.items()):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        set_cell_shading(cell, "1E293B")
        set_cell_borders(cell, color_hex="334155", sz="4")

        # Stat Value
        p_val = cell.paragraphs[0]
        p_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_val.paragraph_format.space_before = Pt(8)
        p_val.paragraph_format.space_after = Pt(2)

        val_color = (239, 68, 68) if i == 2 and val > 0 else (219, 39, 119) if i == 3 and val > 0 else (248, 250, 252)
        format_run(p_val.add_run(str(val)), font_name="Segoe UI", size_pt=20, bold=True, color_rgb=val_color)

        # Stat Label
        p_lbl = cell.add_paragraph()
        p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lbl.paragraph_format.space_before = Pt(0)
        p_lbl.paragraph_format.space_after = Pt(8)
        format_run(p_lbl.add_run(label), font_name="Segoe UI", size_pt=8, bold=True, color_rgb=(148, 163, 184))

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Executive Severity Chart
    chart_buf = generate_severity_chart(vulnerabilities, return_bytes=True)
    if chart_buf:
        chart_p = doc.add_paragraph()
        chart_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        chart_p.paragraph_format.space_after = Pt(15)
        chart_run = chart_p.add_run()
        chart_run.add_picture(chart_buf, width=Inches(5.5))

    doc.add_page_break()

    # Detailed Findings Header
    h_findings = doc.add_paragraph()
    h_findings.paragraph_format.space_after = Pt(12)
    format_run(
        h_findings.add_run("Detailed Findings"), font_name="Segoe UI", size_pt=18, bold=True, color_rgb=(248, 250, 252)
    )

    def get_score(v):
        s = v.get("cvss_score", 0)
        return float(s) if s != "N/A" else 0

    sorted_cves = sorted(vulnerabilities.keys(), key=lambda cid: get_score(vulnerabilities[cid]), reverse=True)

    # Color Palette mapping for Severity
    sev_colors = {
        "Critical": {"bg": "7F1D1D", "text": (252, 165, 165), "name": "CRITICAL"},  # Deep red / pink-red
        "High": {"bg": "7C2D12", "text": (253, 186, 116), "name": "HIGH"},  # Deep orange / light orange
        "Medium": {"bg": "78350F", "text": (253, 224, 71), "name": "MEDIUM"},  # Deep amber / yellow
        "Low": {"bg": "1E3A8A", "text": (147, 197, 253), "name": "LOW"},  # Deep blue / light blue
        "Info": {"bg": "064E3B", "text": (110, 231, 183), "name": "INFO"},  # Deep green / light green
    }

    for i, cve_id in enumerate(sorted_cves):
        d = vulnerabilities[cve_id]
        cvss_val = get_score(d)

        if cvss_val >= 9.0:
            sev_key = "Critical"
        elif cvss_val >= 7.0:
            sev_key = "High"
        elif cvss_val >= 4.0:
            sev_key = "Medium"
        elif cvss_val >= 0.1:
            sev_key = "Low"
        else:
            sev_key = "Info"

        theme = sev_colors[sev_key]

        # -------------------------------------------------------------
        # VULN HEADER CONTAINER (Structured Card Block)
        # -------------------------------------------------------------
        card_table = doc.add_table(rows=1, cols=1)
        card_table.autofit = False
        card_cell = card_table.rows[0].cells[0]
        card_cell.width = Inches(7.0)
        set_cell_shading(card_cell, "1E293B")
        set_cell_borders(card_cell, color_hex="334155", sz="6")

        p_header = card_cell.paragraphs[0]
        p_header.paragraph_format.space_before = Pt(8)
        p_header.paragraph_format.space_after = Pt(2)

        # CVSS Score Badge
        badge_run = p_header.add_run(f" {d.get('cvss_score', 'N/A')} ")
        badge_run.bold = True
        # Emulate a highlighted badge by using highlighted colors if possible,
        # or distinct coloring matching the severity theme
        format_run(badge_run, font_name="Segoe UI", size_pt=13, bold=True, color_rgb=theme["text"])

        # Spacer
        p_header.add_run("   |   ")
        p_header.runs[-1].font.color.rgb = RGBColor(51, 65, 85)

        # CVE ID
        id_run = p_header.add_run(f"{cve_id}   ")
        format_run(id_run, font_name="Segoe UI", size_pt=14, bold=True, color_rgb=(59, 130, 246))

        # Badges (KEV, Exploited)
        if d.get("cisa_kev") == "Yes":
            kev_run = p_header.add_run(" [ CISA KEV ] ")
            format_run(kev_run, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(219, 39, 119))
        if "PoC" in d.get("exploitability", ""):
            poc_run = p_header.add_run(" [ PoC AVAILABLE ] ")
            format_run(poc_run, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(99, 102, 241))
        elif d.get("exploitability") == "Exploited in the wild":
            exp_run = p_header.add_run(" [ ACTIVE EXPLOIT ] ")
            format_run(exp_run, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(220, 38, 38))

        # Title line
        p_title = card_cell.add_paragraph()
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(8)
        format_run(
            p_title.add_run(d.get("title", "N/A")),
            font_name="Segoe UI",
            size_pt=12,
            bold=True,
            color_rgb=(255, 255, 255),
        )

        # Metadata Details Line
        p_meta = card_cell.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(8)

        p_meta.add_run("EPSS: ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        p_meta.add_run(f"{d.get('epss_score', 'N/A')}  |  ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(248, 250, 252))

        p_meta.add_run("Threat Vector: ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        p_meta.add_run(f"{d.get('threat_vector', 'N/A')}  |  ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(248, 250, 252))

        p_meta.add_run("Complexity: ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        p_meta.add_run(f"{d.get('attack_complexity', 'N/A')}")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(248, 250, 252))

        # Add visual separator
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

        # -------------------------------------------------------------
        # TECHNICAL DETAILS
        # -------------------------------------------------------------

        # 1. Description
        p_desc_h = doc.add_paragraph()
        p_desc_h.paragraph_format.space_before = Pt(6)
        p_desc_h.paragraph_format.space_after = Pt(2)
        format_run(
            p_desc_h.add_run("DESCRIPTION"), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246)
        )

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(10)
        format_run(
            p_desc.add_run(d.get("description", "N/A")), font_name="Segoe UI", size_pt=9.5, color_rgb=(203, 213, 225)
        )

        # 2. Affected Products
        p_aff_h = doc.add_paragraph()
        p_aff_h.paragraph_format.space_after = Pt(2)
        format_run(
            p_aff_h.add_run("AFFECTED PRODUCTS"), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246)
        )

        affected_items = d.get("affected", [])
        if affected_items:
            for a in affected_items:
                p_item = doc.add_paragraph(style="List Bullet")
                p_item.paragraph_format.space_after = Pt(1)

                r_prod = p_item.add_run(f"{a['product']} ")
                format_run(r_prod, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(248, 250, 252))

                r_vers = p_item.add_run(f"({', '.join(a['versions'])})")
                format_run(r_vers, font_name="Segoe UI", size_pt=9, color_rgb=(148, 163, 184))
        else:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_after = Pt(4)
            format_run(p_item.add_run("N/A"), font_name="Segoe UI", size_pt=9, italic=True, color_rgb=(148, 163, 184))

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 3. Affected Hosts
        p_hosts_h = doc.add_paragraph()
        p_hosts_h.paragraph_format.space_after = Pt(4)
        format_run(
            p_hosts_h.add_run("AFFECTED HOSTS"), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246)
        )

        hosts = cve_to_hosts.get(cve_id, [])
        if hosts:
            # Generate a nice host tags format inside the document
            p_host_tags = doc.add_paragraph()
            p_host_tags.paragraph_format.space_after = Pt(10)

            # Show up to 5 hosts, matching the HTML limit
            displayed_hosts = hosts[:5]
            for idx, h in enumerate(displayed_hosts):
                host_label = f"{h[1]} ({h[0]})" if h[1] and h[0] else h[1] or h[0]
                tag_text = f" {host_label} [{h[2]}] "

                tag_run = p_host_tags.add_run(tag_text)
                format_run(tag_run, font_name="Consolas", size_pt=8.5, color_rgb=(148, 163, 184))

                # Add inline visual tag shading if possible, or simple coloring
                if idx < len(displayed_hosts) - 1:
                    sep = p_host_tags.add_run("   ")
                    format_run(sep, font_name="Segoe UI")

            if len(hosts) > 5:
                more_run = p_host_tags.add_run(f"   (+{len(hosts)-5} more)")
                format_run(more_run, font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(59, 130, 246))
        else:
            p_none = doc.add_paragraph()
            p_none.paragraph_format.space_after = Pt(10)
            format_run(p_none.add_run("N/A"), font_name="Segoe UI", size_pt=9, italic=True, color_rgb=(148, 163, 184))

        # 4. Remediation Recommendation Card
        rem_table = doc.add_table(rows=1, cols=1)
        rem_table.autofit = False
        rem_cell = rem_table.rows[0].cells[0]
        rem_cell.width = Inches(7.0)

        # Shade with a deep blue tint block to match HTML (.remediation)
        set_cell_shading(rem_cell, "0F1E36")  # Very deep blue
        set_cell_borders(rem_cell, color_hex="1D4ED8", sz="4")  # Vibrant blue border

        p_rem_title = rem_cell.paragraphs[0]
        p_rem_title.paragraph_format.space_before = Pt(6)
        p_rem_title.paragraph_format.space_after = Pt(2)
        format_run(
            p_rem_title.add_run("🛡️ Remediation Recommendation"),
            font_name="Segoe UI",
            size_pt=8.5,
            bold=True,
            color_rgb=(59, 130, 246),
        )

        p_rem_val = rem_cell.add_paragraph()
        p_rem_val.paragraph_format.space_after = Pt(6)
        format_run(
            p_rem_val.add_run(d.get("remediation", "N/A")),
            font_name="Segoe UI",
            size_pt=9.5,
            bold=True,
            color_rgb=(248, 250, 252),
        )

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 5. Intelligence & References
        p_ref_h = doc.add_paragraph()
        p_ref_h.paragraph_format.space_after = Pt(2)
        format_run(
            p_ref_h.add_run("INTELLIGENCE & REFERENCES"),
            font_name="Segoe UI",
            size_pt=9.5,
            bold=True,
            color_rgb=(59, 130, 246),
        )

        if d.get("poc_link") != "N/A":
            p_poc = doc.add_paragraph(style="List Bullet")
            p_poc.paragraph_format.space_after = Pt(1)
            format_run(
                p_poc.add_run("🚀 Exploit PoC Link: "),
                font_name="Segoe UI",
                size_pt=9,
                bold=True,
                color_rgb=(248, 250, 252),
            )
            poc_link_run = p_poc.add_run(d["poc_link"])
            format_run(poc_link_run, font_name="Segoe UI", size_pt=9, color_rgb=(59, 130, 246))

        for ref in d.get("references", [])[:4]:
            p_ref = doc.add_paragraph(style="List Bullet")
            p_ref.paragraph_format.space_after = Pt(1)
            ref_run = p_ref.add_run(ref)
            format_run(ref_run, font_name="Segoe UI", size_pt=9, color_rgb=(59, 130, 246))

        # Page break if not the last item
        if i < len(sorted_cves) - 1:
            doc.add_page_break()

    # Save to output file
    doc.save(output_file)
    print(f"[+] Word (DOCX) intelligence report saved to: {output_file}")
