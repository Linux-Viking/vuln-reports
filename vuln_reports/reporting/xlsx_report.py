"""Excel workbook: a summary sheet plus one sheet per CVE with its affected hosts."""

from vuln_reports.reporting.common import get_score

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def save_xlsx_report(vulnerabilities, cve_to_hosts, output_file):
    if not HAS_OPENPYXL:
        print("[!] openpyxl not installed. XLSX report skipped.")
        return

    wb = openpyxl.Workbook()

    # 1. Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    summary_headers = ["CVE ID", "Title", "CVSS Score", "CISA KEV", "Exploitability", "Hosts Affected"]
    ws_summary.append(summary_headers)

    header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    sorted_cves = sorted(vulnerabilities.keys(), key=lambda cid: get_score(vulnerabilities[cid]), reverse=True)

    for cve_id in sorted_cves:
        d = vulnerabilities[cve_id]
        ws_summary.append(
            [
                cve_id,
                d.get("title", "N/A"),
                d.get("cvss_score", "N/A"),
                d.get("cisa_kev", "N/A"),
                d.get("exploitability", "N/A"),
                len(cve_to_hosts.get(cve_id, [])),
            ]
        )

    # 2. One sheet per CVE (Excel sheet names are limited to 31 chars)
    for cve_id in sorted_cves:
        d = vulnerabilities[cve_id]
        ws = wb.create_sheet(title=cve_id[:31])

        metadata = [
            ("CVE ID", cve_id),
            ("Title", d.get("title", "N/A")),
            ("Description", d.get("description", "N/A")),
            ("CVSS Score", d.get("cvss_score", "N/A")),
            ("CISA KEV", d.get("cisa_kev", "N/A")),
            ("Exploitability", d.get("exploitability", "N/A")),
            ("EPSS Score", d.get("epss_score", "N/A")),
            ("Remediation", d.get("remediation", "N/A")),
        ]
        for i, (label, val) in enumerate(metadata, 1):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=str(val)).alignment = Alignment(wrap_text=True)

        start_row = len(metadata) + 2
        ws.cell(row=start_row, column=1, value="Affected Hosts").font = Font(bold=True, size=12)
        for col, h_text in enumerate(["IP Address", "Hostname", "Port/Proto", "Service Name"], 1):
            cell = ws.cell(row=start_row + 1, column=col, value=h_text)
            cell.fill = header_fill
            cell.font = header_font

        for ip, host, port, svc in sorted(cve_to_hosts.get(cve_id, [])):
            ws.append([ip, host, port, svc])

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 80

    wb.save(output_file)
    print(f"[+] XLSX intelligence report saved to: {output_file}")
