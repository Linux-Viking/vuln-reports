import csv
import json
import xml.etree.ElementTree as ET

import pytest

from vuln_reports.models import EnrichedScan, Host, ScanResult, Service
from vuln_reports.reporting import chart, common, lookup_formats
from vuln_reports.reporting.csv_report import HEADERS, save_scan_csv
from vuln_reports.reporting.docx_report import save_docx_report
from vuln_reports.reporting.html_report import save_html_report, save_markdown_report
from vuln_reports.reporting.xlsx_report import save_xlsx_report

INTEL = {
    "cve_id": "CVE-2021-0001",
    "title": "=HYPERLINK(evil)",
    "description": "Line one\nline | two",
    "affected": [{"vendor": "V", "product": "P", "versions": ["1.0", "2.0"]}],
    "cisa_kev": "Yes",
    "cvss_score": 9.8,
    "exploitability": "Exploited in the wild",
    "poc_available": "Yes",
    "poc_link": "https://example.test/poc",
    "threat_vector": "NETWORK",
    "user_interaction": "NONE",
    "epss_score": "0.5",
    "attack_complexity": "LOW",
    "references": ["https://example.test/a", "https://example.test/b"],
    "remediation": "Patch",
}


@pytest.mark.parametrize(
    "score, bucket, level",
    [
        (9.0, "Critical", "critical"),
        (7.0, "High", "high"),
        (4.0, "Medium", "medium"),
        (0.1, "Low", "low"),
        (0.0, "Info", "info"),
    ],
)
def test_severity_mappings(score, bucket, level):
    assert common.severity_bucket(score) == bucket
    assert common.get_severity(score) == level


def test_get_severity_edge_cases():
    assert common.get_severity("N/A") == "info"
    assert common.get_severity("bogus") == "info"
    assert common.get_severity(0.05) == "low"  # filter treats >0 as Low; the chart starts Low at 0.1


def test_compute_stats_and_sorting():
    vulns = {"A": {"cvss_score": 5.0}, "B": {"cvss_score": 9.5, "cisa_kev": "Yes"}, "C": {"cvss_score": "N/A"}}
    hosts = {
        "A": [("192.0.2.1", "", "22/tcp", "ssh")],
        "B": [("192.0.2.1", "", "80/tcp", "http"), ("192.0.2.2", "", "80/tcp", "http")],
    }
    stats = common.compute_stats(vulns, hosts)
    assert stats == {
        "total_hosts": 2,
        "total_cves": 3,
        "critical_count": 1,
        "high_count": 0,
        "medium_count": 1,
        "low_count": 0,
        "info_count": 1,
        "kev_count": 1,
    }
    assert list(common.sorted_by_score(vulns)) == ["B", "A", "C"]


def test_chart_counts_and_empty():
    assert chart.severity_counts({"a": {"cvss_score": "7.5"}, "b": {}}) == {
        "Critical": 0,
        "High": 1,
        "Medium": 0,
        "Low": 0,
        "Info": 1,
    }
    assert chart.generate_severity_chart({}) is None


def test_csv_safe():
    assert lookup_formats.csv_safe("=1+1") == "'=1+1"
    assert lookup_formats.csv_safe("-x") == "'-x"
    assert lookup_formats.csv_safe("ok") == "ok"


def test_get_output_path_never_overwrites(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert lookup_formats.get_output_path(None, "csv") == "results.csv"
    assert lookup_formats.get_output_path("out.csv", "csv") == "out.csv"
    (tmp_path / "out.csv").write_text("x")
    (tmp_path / "out_1.csv").write_text("x")
    assert lookup_formats.get_output_path("out", "csv") == "out_2.csv"


def test_lookup_csv_contract(tmp_path):
    path = tmp_path / "r.csv"
    lookup_formats.save_csv([INTEL, {"cve_id": "CVE-2099-1", "error": "x"}], str(path))
    rows = list(csv.reader(path.open(newline="", encoding="utf-8")))
    assert rows[0] == lookup_formats.CSV_HEADERS
    assert len(rows) == 2  # error rows are skipped
    assert rows[1][1] == "'=HYPERLINK(evil)"
    assert rows[1][11] == "P (1.0, 2.0)"


def test_grep_xml_text_json(tmp_path):
    lookup_formats.save_grep([INTEL], str(tmp_path / "r.grep"))
    line = (tmp_path / "r.grep").read_text().rstrip("\n")
    fields = line.split("|")
    assert len(fields) == 14 and fields[2] == "Line one line   two"

    lookup_formats.save_xml([INTEL], str(tmp_path / "r.xml"))
    assert ET.parse(tmp_path / "r.xml").getroot().find("CVE/cve_id").text == "CVE-2021-0001"

    lookup_formats.save_text([INTEL], str(tmp_path / "r.txt"))
    assert "Affected: P (1.0, 2.0)" in (tmp_path / "r.txt").read_text()

    lookup_formats.save_json([INTEL], str(tmp_path / "r.json"))
    assert json.loads((tmp_path / "r.json").read_text()) == [INTEL]


def test_print_pretty_minimal_hides_missing_poc(capsys):
    lookup_formats.print_pretty({**INTEL, "poc_link": "N/A"}, use_colors=False, minimal=True)
    assert capsys.readouterr().out == ""
    lookup_formats.print_pretty(INTEL, use_colors=False, minimal=True)
    assert capsys.readouterr().out == "CVE-2021-0001: https://example.test/poc\n"


def make_enriched():
    scan = ScanResult(
        "nessus_xml",
        "/scans/in.nessus",
        {"192.0.2.1": Host("192.0.2.1", "a.example.test", [Service("443", "tcp", "www", {"CVE-2021-0001"})])},
    )
    return EnrichedScan(scan, {"CVE-2021-0001": INTEL})


def test_scan_csv(tmp_path):
    path = tmp_path / "scan.csv"
    save_scan_csv(make_enriched(), str(path))
    rows = list(csv.reader(path.open(newline="", encoding="utf-8")))
    assert rows[0] == HEADERS
    assert rows[1][:5] == ["192.0.2.1", "a.example.test", "CVE-2021-0001", "443/tcp", "www"]


def test_html_markdown_xlsx_docx(tmp_path):
    enriched = make_enriched()
    hosts = enriched.scan.cve_to_hosts()
    save_html_report(enriched.intel, hosts, "nessus_xml", "/scans/in.nessus", str(tmp_path / "r.html"))
    html = (tmp_path / "r.html").read_text(encoding="utf-8")
    assert "CVE-2021-0001" in html and "in.nessus" in html and "data:image/png;base64," in html

    save_markdown_report(enriched.intel, hosts, "nessus_xml", "/scans/in.nessus", str(tmp_path / "r.md"))
    assert "## [CVE-2021-0001]" in (tmp_path / "r.md").read_text(encoding="utf-8")

    import openpyxl

    save_xlsx_report(enriched.intel, hosts, str(tmp_path / "r.xlsx"))
    wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
    assert wb.sheetnames == ["Summary", "CVE-2021-0001"]
    assert wb["Summary"]["F2"].value == 1

    import docx

    save_docx_report(enriched.intel, hosts, "nessus_xml", "/scans/in.nessus", str(tmp_path / "r.docx"))
    text = "\n".join(p.text for p in docx.Document(tmp_path / "r.docx").paragraphs)
    assert "Detailed Findings" in text and "AFFECTED HOSTS" in text
