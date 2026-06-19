#!/usr/bin/env python3
"""
VulnReport: Unified vulnerability report generator.
Combines Scan2CVE (parsing) and CVE-Lookup (enrichment) to produce a detailed CSV.
"""

import sys
import os

# --- Aggressive GLib/GIO Warning Suppression ---
# These must be set BEFORE any GIO/GTK-based libraries are loaded.
os.environ["GIO_USE_VFS"] = "local"
os.environ["GIO_USE_DESKTOP_APP_INFO"] = "none"
os.environ["GIO_USE_VOLUME_MONITOR"] = "none"
os.environ["G_MESSAGES_DEBUG"] = "none"

# Function to suppress C-level stderr (to catch persistent GLib warnings on Windows)
def suppress_gio_noise():
    if os.name != 'nt': return None
    try:
        null_fd = os.open(os.devnull, os.O_RDWR)
        old_stderr = os.dup(sys.stderr.fileno())
        os.dup2(null_fd, sys.stderr.fileno())
        os.close(null_fd)
        return old_stderr
    except: return None

def restore_gio_noise(old_stderr):
    if old_stderr is not None:
        os.dup2(old_stderr, sys.stderr.fileno())
        os.close(old_stderr)

# Suppress noise during initialization imports
_gio_stderr = suppress_gio_noise()
try:
    import csv
    import argparse
    import requests
    import time
    from typing import List, Dict, Any

    # Dynamic imports
    import scan2cve
    import cve_lookup

    from concurrent.futures import ThreadPoolExecutor, as_completed
    from datetime import datetime

    try:
        from jinja2 import Environment, FileSystemLoader
        HAS_JINJA = True
    except Exception:
        HAS_JINJA = False

    try:
        from weasyprint import HTML
        HAS_WEASYPRINT = True
    except Exception as e:
        HAS_WEASYPRINT = False

    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
        HAS_OPENPYXL = True
    except ImportError:
        HAS_OPENPYXL = False

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        HAS_DOCX = True
    except ImportError:
        HAS_DOCX = False

    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        import io
        import base64
        HAS_MATPLOTLIB = True
    except ImportError:
        HAS_MATPLOTLIB = False
finally:
    restore_gio_noise(_gio_stderr)

def get_detailed_cve_info(cve_ids: List[str], github_token: str = None, nvd_key: str = None, vulners_key: str = None, circl_key: str = None, max_workers: int = 10) -> Dict[str, Dict[str, Any]]:
    """Fetch details for a list of CVEs using cve_lookup logic with concurrency."""
    session = requests.Session()
    cve_details = {}
    print(f"[*] Fetching enrichment data for {len(cve_ids)} unique CVEs...")

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(cve_lookup.get_cve_data, session, cve_id, github_token, nvd_key, vulners_key, circl_key): cve_id for cve_id in cve_ids}
        for future in as_completed(futures):
            cve_id = futures[future]
            try:
                data = future.result()
                if "error" not in data:
                    cve_details[cve_id] = data
                else:
                    print(f"  [!] Error enriching {cve_id}: {data['error']}")
                    cve_details[cve_id] = {"cve_id": cve_id, "error": data["error"]}
            except Exception as e:
                print(f"  [!] Exception enriching {cve_id}: {e}")
                cve_details[cve_id] = {"cve_id": cve_id, "error": str(e)}
    return cve_details

def generate_severity_chart(vulnerabilities: Dict[str, Any], return_bytes: bool = False):
    """Generates a severity distribution pie chart and returns it as a base64 string or bytes."""
    if not HAS_MATPLOTLIB:
        return None

    counts = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0, "Info": 0}
    
    for v in vulnerabilities.values():
        score = v.get('cvss_score')
        if not isinstance(score, (int, float)):
            try:
                if score == "N/A" or score is None:
                    score = 0
                else:
                    score = float(score)
            except (ValueError, TypeError):
                score = 0
        
        if score >= 9.0:
            counts["Critical"] += 1
        elif score >= 7.0:
            counts["High"] += 1
        elif score >= 4.0:
            counts["Medium"] += 1
        elif score >= 0.1:
            counts["Low"] += 1
        else:
            counts["Info"] += 1

    labels = [k for k in counts.keys() if counts[k] > 0]
    sizes = [counts[k] for k in counts.keys() if counts[k] > 0]
    
    if not sizes:
        return None

    colors = {
        "Critical": "#dc2626", # red-600
        "High": "#ea580c",     # orange-600
        "Medium": "#f59e0b",   # amber-500
        "Low": "#3b82f6",      # blue-500
        "Info": "#22c55e"      # green-500
    }
    chart_colors = [colors[l] for l in labels]

    # Create figure with dark background to match report theme
    plt.figure(figsize=(7, 5), facecolor='#0F172A')
    plt.rcParams['text.color'] = 'white'
    
    patches, texts, autotexts = plt.pie(
        sizes, 
        labels=labels, 
        autopct='%1.1f%%', 
        startangle=140, 
        colors=chart_colors,
        textprops={'fontsize': 10, 'weight': 'bold'}
    )
    
    for text in texts:
        text.set_color('white')
    for autotext in autotexts:
        autotext.set_color('white')

    plt.title("Severity Distribution", color='white', fontsize=14, fontweight='bold', pad=10)
    plt.axis('equal')
    
    buf = io.BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight', facecolor='#0F172A', dpi=150)
    plt.close()
    buf.seek(0)
    
    if return_bytes:
        return buf
    return base64.b64encode(buf.read()).decode('utf-8')

def save_docx_report(vulnerabilities, cve_to_hosts, scanner_type, target_file, output_file):
    if not HAS_DOCX:
        print("[!] python-docx not installed. DOCX report skipped.")
        return

    from docx.oxml.ns import nsdecls
    from docx.oxml import parse_xml
    from docx.oxml.shared import qn

    # --- TEMPLATE LOGIC ---
    template_path = os.path.join(os.path.dirname(__file__), 'templates', 'report_template.docx')
    if os.path.exists(template_path):
        doc = Document(template_path)
        print(f"[*] Using custom Word template: {template_path}")
    else:
        doc = Document()
    
    # Configure page margins for a cleaner dashboard look (0.75 in margins)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # --- GLOBAL DARK THEME ---
    # 1. Set Page Background Color (#0F172A)
    shading_elm_2 = parse_xml(r'<w:background {} w:color="0F172A"/>'.format(nsdecls('w')))
    doc.element.insert(0, shading_elm_2)
    
    # 2. Enable background color display in Word settings
    settings = doc.settings.element
    display_bg = parse_xml(r'<w:displayBackgroundShape {}/>'.format(nsdecls('w')))
    settings.append(display_bg)

    # Helper to set cell shading (background color)
    def set_cell_shading(cell, color_hex):
        shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
        cell._tc.get_or_add_tcPr().append(shading)

    # Helper to set cell borders
    def set_cell_borders(cell, color_hex="334155", sz="4", val="single"):
        borders = parse_xml(
            r'<w:tcBorders {}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>'
            .format(nsdecls('w'), val=val, sz=sz, color=color_hex)
        )
        cell._tc.get_or_add_tcPr().append(borders)

    # Helper to format fonts and colors
    def format_run(run, font_name="Segoe UI", size_pt=10, bold=False, italic=False, color_rgb=(248, 250, 252)):
        run.font.name = font_name
        # Apply ASCII/Hansi/CS font settings in oxml to prevent fallback
        rPr = run._r.get_or_add_rPr()
        rFonts = parse_xml(r'<w:rFonts {} w:ascii="{f}" w:hAnsi="{f}" w:cs="{f}"/>'.format(nsdecls('w'), f=font_name))
        rPr.append(rFonts)
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)
        return run

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    format_run(title_p.add_run('Vulnerability Intelligence Report'), font_name="Segoe UI", size_pt=24, bold=True, color_rgb=(59, 130, 246))
    
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    format_run(meta_p.add_run(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Scanner: {scanner_type} | Target: {os.path.basename(target_file)}"), font_name="Segoe UI", size_pt=9.5, italic=True, color_rgb=(148, 163, 184))

    # Summary Stats Section
    h_exec = doc.add_paragraph()
    h_exec.paragraph_format.space_after = Pt(8)
    format_run(h_exec.add_run('Executive Summary'), font_name="Segoe UI", size_pt=16, bold=True, color_rgb=(248, 250, 252))

    stats = {
        "HOSTS": len(set(h[0] for hosts in cve_to_hosts.values() for h in hosts)),
        "Unique CVEs": len(vulnerabilities),
        "Critical Risks": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and v['cvss_score'] >= 9),
        "KEV Exploited": sum(1 for v in vulnerabilities.values() if v.get('cisa_kev') == 'Yes')
    }

    # Generate 4 distinct columns for cards
    stats_table = doc.add_table(rows=1, cols=4)
    stats_table.autofit = False
    col_widths = [Inches(1.75), Inches(1.75), Inches(1.75), Inches(1.75)]
    
    hdr_cells = stats_table.rows[0].cells
    for i, (label, val) in enumerate(stats.items()):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        set_cell_shading(cell, "1E293B")
        set_cell_borders(cell, color_hex="334155", sz="4")
        
        # Stat Value
        p_val = cell.paragraphs[0]
        p_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_val.paragraph_format.space_before = Pt(8)
        p_val.paragraph_format.space_after = Pt(2)
        
        val_color = (239, 68, 68) if i == 2 and val > 0 else (219, 39, 119) if i == 3 and val > 0 else (248, 250, 252)
        format_run(p_val.add_run(str(val)), font_name="Segoe UI", size_pt=20, bold=True, color_rgb=val_color)
        
        # Stat Label
        p_lbl = cell.add_paragraph()
        p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lbl.paragraph_format.space_before = Pt(0)
        p_lbl.paragraph_format.space_after = Pt(8)
        format_run(p_lbl.add_run(label), font_name="Segoe UI", size_pt=8, bold=True, color_rgb=(148, 163, 184))

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Executive Severity Chart
    chart_buf = generate_severity_chart(vulnerabilities, return_bytes=True)
    if chart_buf:
        chart_p = doc.add_paragraph()
        chart_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        chart_p.paragraph_format.space_after = Pt(15)
        chart_run = chart_p.add_run()
        chart_run.add_picture(chart_buf, width=Inches(5.5))
        
    doc.add_page_break()

    # Detailed Findings Header
    h_findings = doc.add_paragraph()
    h_findings.paragraph_format.space_after = Pt(12)
    format_run(h_findings.add_run('Detailed Findings'), font_name="Segoe UI", size_pt=18, bold=True, color_rgb=(248, 250, 252))

    def get_score(v):
        s = v.get('cvss_score', 0)
        return float(s) if s != 'N/A' else 0

    sorted_cves = sorted(vulnerabilities.keys(), key=lambda cid: get_score(vulnerabilities[cid]), reverse=True)

    # Color Palette mapping for Severity
    sev_colors = {
        "Critical": {"bg": "7F1D1D", "text": (252, 165, 165), "name": "CRITICAL"}, # Deep red / pink-red
        "High": {"bg": "7C2D12", "text": (253, 186, 116), "name": "HIGH"},       # Deep orange / light orange
        "Medium": {"bg": "78350F", "text": (253, 224, 71), "name": "MEDIUM"},    # Deep amber / yellow
        "Low": {"bg": "1E3A8A", "text": (147, 197, 253), "name": "LOW"},         # Deep blue / light blue
        "Info": {"bg": "064E3B", "text": (110, 231, 183), "name": "INFO"}        # Deep green / light green
    }

    for i, cve_id in enumerate(sorted_cves):
        d = vulnerabilities[cve_id]
        cvss_val = get_score(d)
        
        if cvss_val >= 9.0: sev_key = "Critical"
        elif cvss_val >= 7.0: sev_key = "High"
        elif cvss_val >= 4.0: sev_key = "Medium"
        elif cvss_val >= 0.1: sev_key = "Low"
        else: sev_key = "Info"
        
        theme = sev_colors[sev_key]

        # -------------------------------------------------------------
        # VULN HEADER CONTAINER (Structured Card Block)
        # -------------------------------------------------------------
        card_table = doc.add_table(rows=1, cols=1)
        card_table.autofit = False
        card_cell = card_table.rows[0].cells[0]
        card_cell.width = Inches(7.0)
        set_cell_shading(card_cell, "1E293B")
        set_cell_borders(card_cell, color_hex="334155", sz="6")

        p_header = card_cell.paragraphs[0]
        p_header.paragraph_format.space_before = Pt(8)
        p_header.paragraph_format.space_after = Pt(2)
        
        # CVSS Score Badge
        badge_run = p_header.add_run(f" {d.get('cvss_score', 'N/A')} ")
        badge_run.bold = True
        # Emulate a highlighted badge by using highlighted colors if possible, 
        # or distinct coloring matching the severity theme
        format_run(badge_run, font_name="Segoe UI", size_pt=13, bold=True, color_rgb=theme["text"])
        
        # Spacer
        p_header.add_run("   |   ")
        p_header.runs[-1].font.color.rgb = RGBColor(51, 65, 85)
        
        # CVE ID
        id_run = p_header.add_run(f"{cve_id}   ")
        format_run(id_run, font_name="Segoe UI", size_pt=14, bold=True, color_rgb=(59, 130, 246))

        # Badges (KEV, Exploited)
        if d.get('cisa_kev') == 'Yes':
            kev_run = p_header.add_run(" [ CISA KEV ] ")
            format_run(kev_run, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(219, 39, 119))
        if "PoC" in d.get('exploitability', ''):
            poc_run = p_header.add_run(" [ PoC AVAILABLE ] ")
            format_run(poc_run, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(99, 102, 241))
        elif d.get('exploitability') == "Exploited in the wild":
            exp_run = p_header.add_run(" [ ACTIVE EXPLOIT ] ")
            format_run(exp_run, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(220, 38, 38))

        # Title line
        p_title = card_cell.add_paragraph()
        p_title.paragraph_format.space_before = Pt(4)
        p_title.paragraph_format.space_after = Pt(8)
        format_run(p_title.add_run(d.get('title', 'N/A')), font_name="Segoe UI", size_pt=12, bold=True, color_rgb=(255, 255, 255))

        # Metadata Details Line
        p_meta = card_cell.add_paragraph()
        p_meta.paragraph_format.space_after = Pt(8)
        
        p_meta.add_run("EPSS: ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        p_meta.add_run(f"{d.get('epss_score', 'N/A')}  |  ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(248, 250, 252))

        p_meta.add_run("Threat Vector: ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        p_meta.add_run(f"{d.get('threat_vector', 'N/A')}  |  ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(248, 250, 252))

        p_meta.add_run("Complexity: ")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, color_rgb=(148, 163, 184))
        p_meta.add_run(f"{d.get('attack_complexity', 'N/A')}")
        format_run(p_meta.runs[-1], font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(248, 250, 252))

        # Add visual separator
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

        # -------------------------------------------------------------
        # TECHNICAL DETAILS
        # -------------------------------------------------------------
        
        # 1. Description
        p_desc_h = doc.add_paragraph()
        p_desc_h.paragraph_format.space_before = Pt(6)
        p_desc_h.paragraph_format.space_after = Pt(2)
        format_run(p_desc_h.add_run('DESCRIPTION'), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246))
        
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(10)
        format_run(p_desc.add_run(d.get('description', 'N/A')), font_name="Segoe UI", size_pt=9.5, color_rgb=(203, 213, 225))

        # 2. Affected Products
        p_aff_h = doc.add_paragraph()
        p_aff_h.paragraph_format.space_after = Pt(2)
        format_run(p_aff_h.add_run('AFFECTED PRODUCTS'), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246))
        
        affected_items = d.get('affected', [])
        if affected_items:
            for a in affected_items:
                p_item = doc.add_paragraph(style='List Bullet')
                p_item.paragraph_format.space_after = Pt(1)
                
                r_prod = p_item.add_run(f"{a['product']} ")
                format_run(r_prod, font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(248, 250, 252))
                
                r_vers = p_item.add_run(f"({', '.join(a['versions'])})")
                format_run(r_vers, font_name="Segoe UI", size_pt=9, color_rgb=(148, 163, 184))
        else:
            p_item = doc.add_paragraph()
            p_item.paragraph_format.space_after = Pt(4)
            format_run(p_item.add_run("N/A"), font_name="Segoe UI", size_pt=9, italic=True, color_rgb=(148, 163, 184))

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 3. Affected Hosts
        p_hosts_h = doc.add_paragraph()
        p_hosts_h.paragraph_format.space_after = Pt(4)
        format_run(p_hosts_h.add_run('AFFECTED HOSTS'), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246))

        hosts = cve_to_hosts.get(cve_id, [])
        if hosts:
            # Generate a nice host tags format inside the document
            p_host_tags = doc.add_paragraph()
            p_host_tags.paragraph_format.space_after = Pt(10)
            
            # Show up to 5 hosts, matching the HTML limit
            displayed_hosts = hosts[:5]
            for idx, h in enumerate(displayed_hosts):
                host_label = f"{h[1]} ({h[0]})" if h[1] and h[0] else h[1] or h[0]
                tag_text = f" {host_label} [{h[2]}] "
                
                tag_run = p_host_tags.add_run(tag_text)
                format_run(tag_run, font_name="Consolas", size_pt=8.5, color_rgb=(148, 163, 184))
                
                # Add inline visual tag shading if possible, or simple coloring
                if idx < len(displayed_hosts) - 1:
                    sep = p_host_tags.add_run("   ")
                    format_run(sep, font_name="Segoe UI")
            
            if len(hosts) > 5:
                more_run = p_host_tags.add_run(f"   (+{len(hosts)-5} more)")
                format_run(more_run, font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(59, 130, 246))
        else:
            p_none = doc.add_paragraph()
            p_none.paragraph_format.space_after = Pt(10)
            format_run(p_none.add_run("N/A"), font_name="Segoe UI", size_pt=9, italic=True, color_rgb=(148, 163, 184))

        # 4. Remediation Recommendation Card
        rem_table = doc.add_table(rows=1, cols=1)
        rem_table.autofit = False
        rem_cell = rem_table.rows[0].cells[0]
        rem_cell.width = Inches(7.0)
        
        # Shade with a deep blue tint block to match HTML (.remediation)
        set_cell_shading(rem_cell, "0F1E36") # Very deep blue
        set_cell_borders(rem_cell, color_hex="1D4ED8", sz="4") # Vibrant blue border
        
        p_rem_title = rem_cell.paragraphs[0]
        p_rem_title.paragraph_format.space_before = Pt(6)
        p_rem_title.paragraph_format.space_after = Pt(2)
        format_run(p_rem_title.add_run("🛡️ Remediation Recommendation"), font_name="Segoe UI", size_pt=8.5, bold=True, color_rgb=(59, 130, 246))
        
        p_rem_val = rem_cell.add_paragraph()
        p_rem_val.paragraph_format.space_after = Pt(6)
        format_run(p_rem_val.add_run(d.get('remediation', 'N/A')), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(248, 250, 252))

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 5. Intelligence & References
        p_ref_h = doc.add_paragraph()
        p_ref_h.paragraph_format.space_after = Pt(2)
        format_run(p_ref_h.add_run('INTELLIGENCE & REFERENCES'), font_name="Segoe UI", size_pt=9.5, bold=True, color_rgb=(59, 130, 246))

        if d.get('poc_link') != 'N/A':
            p_poc = doc.add_paragraph(style='List Bullet')
            p_poc.paragraph_format.space_after = Pt(1)
            format_run(p_poc.add_run("🚀 Exploit PoC Link: "), font_name="Segoe UI", size_pt=9, bold=True, color_rgb=(248, 250, 252))
            poc_link_run = p_poc.add_run(d['poc_link'])
            format_run(poc_link_run, font_name="Segoe UI", size_pt=9, color_rgb=(59, 130, 246))

        for ref in d.get('references', [])[:4]:
            p_ref = doc.add_paragraph(style='List Bullet')
            p_ref.paragraph_format.space_after = Pt(1)
            ref_run = p_ref.add_run(ref)
            format_run(ref_run, font_name="Segoe UI", size_pt=9, color_rgb=(59, 130, 246))

        # Page break if not the last item
        if i < len(sorted_cves) - 1:
            doc.add_page_break()

    # Save to output file
    doc.save(output_file)
    print(f"[+] Word (DOCX) intelligence report saved to: {output_file}")

def save_xlsx_report(vulnerabilities, cve_to_hosts, output_file):
    if not HAS_OPENPYXL:
        print("[!] openpyxl not installed. XLSX report skipped.")
        return

    wb = openpyxl.Workbook()
    
    # 1. Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    summary_headers = ["CVE ID", "Title", "CVSS Score", "CISA KEV", "Exploitability", "Hosts Affected"]
    ws_summary.append(summary_headers)
    
    # Header Styling
    header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    # Sort vulnerabilities by CVSS
    def get_score(v):
        s = v.get('cvss_score', 0)
        return float(s) if s != 'N/A' else 0
    
    sorted_cves = sorted(vulnerabilities.keys(), key=lambda cid: get_score(vulnerabilities[cid]), reverse=True)

    for cve_id in sorted_cves:
        d = vulnerabilities[cve_id]
        hosts_count = len(cve_to_hosts.get(cve_id, []))
        ws_summary.append([
            cve_id, d.get('title', 'N/A'), d.get('cvss_score', 'N/A'),
            d.get('cisa_kev', 'N/A'), d.get('exploitability', 'N/A'), hosts_count
        ])

    # 2. Individual CVE Sheets
    for cve_id in sorted_cves:
        d = vulnerabilities[cve_id]
        # Excel sheet names limited to 31 chars
        sheet_name = cve_id[:31]
        ws = wb.create_sheet(title=sheet_name)
        
        # CVE Metadata Table
        metadata = [
            ("CVE ID", cve_id),
            ("Title", d.get('title', 'N/A')),
            ("Description", d.get('description', 'N/A')),
            ("CVSS Score", d.get('cvss_score', 'N/A')),
            ("CISA KEV", d.get('cisa_kev', 'N/A')),
            ("Exploitability", d.get('exploitability', 'N/A')),
            ("EPSS Score", d.get('epss_score', 'N/A')),
            ("Remediation", d.get('remediation', 'N/A'))
        ]
        
        for i, (label, val) in enumerate(metadata, 1):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=str(val)).alignment = Alignment(wrap_text=True)
        
        # Affected Hosts Table
        start_row = len(metadata) + 2
        ws.cell(row=start_row, column=1, value="Affected Hosts").font = Font(bold=True, size=12)
        host_headers = ["IP Address", "Hostname", "Port/Proto", "Service Name"]
        for col, h_text in enumerate(host_headers, 1):
            cell = ws.cell(row=start_row+1, column=col, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            
        for i, (ip, host, port, svc) in enumerate(sorted(cve_to_hosts.get(cve_id, [])), 1):
            ws.append([ip, host, port, svc]) 

        # Adjust columns
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 80

    wb.save(output_file)
    print(f"[+] XLSX intelligence report saved to: {output_file}")

def save_markdown_report(vulnerabilities, cve_to_hosts, scanner_type, target_file, output_file):
    if not HAS_JINJA: return
    
    # Calculate statistics
    stats = {
        "total_hosts": len(set(h[0] for hosts in cve_to_hosts.values() for h in hosts)),
        "total_cves": len(vulnerabilities),
        "critical_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and v['cvss_score'] >= 9),
        "high_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 7 <= v['cvss_score'] < 9),
        "medium_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 4 <= v['cvss_score'] < 7),
        "low_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 0.1 <= v['cvss_score'] < 4),
        "info_count": sum(1 for v in vulnerabilities.values() if not isinstance(v.get('cvss_score'), (int, float)) or v['cvss_score'] < 0.1),
        "kev_count": sum(1 for v in vulnerabilities.values() if v.get('cisa_kev') == 'Yes')
    }

    chart_b64 = generate_severity_chart(vulnerabilities)

    try:
        env = Environment(loader=FileSystemLoader(os.path.join(os.path.dirname(__file__), 'templates')))
        template = env.get_template('report_template.md')

        # Sort vulnerabilities by CVSS score descending
        def get_score(v):
            s = v.get('cvss_score', 0)
            return float(s) if s != 'N/A' else 0

        sorted_vulns = dict(sorted(vulnerabilities.items(), key=lambda x: get_score(x[1]), reverse=True))

        md_out = template.render(
            vulnerabilities=sorted_vulns,
            cve_to_hosts=cve_to_hosts,
            stats=stats,
            chart_b64=chart_b64,
            date=datetime.now().strftime("%Y-%m-%d %H:%M"),
            scanner_type=scanner_type,
            target_file=os.path.basename(target_file)
        )
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(md_out)
        print(f"[+] Markdown report saved to: {output_file}")
    except Exception as e:
        print(f"[!] Error generating Markdown report: {e}")

def save_pdf_report(vulnerabilities, cve_to_hosts, scanner_type, target_file, output_file):
    if not HAS_WEASYPRINT or not HAS_JINJA:
        print("[!] PDF generation requires WeasyPrint and Jinja2.")
        print("[*] Note: On Windows, WeasyPrint also requires the GTK+ runtime installed and in your PATH.")
        print("[!] PDF report skipped.")
        return

    # Generate HTML first (temporarily or in memory)
    stats = {
        "total_hosts": len(set(h[0] for hosts in cve_to_hosts.values() for h in hosts)),
        "total_cves": len(vulnerabilities),
        "critical_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and v['cvss_score'] >= 9),
        "high_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 7 <= v['cvss_score'] < 9),
        "medium_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 4 <= v['cvss_score'] < 7),
        "low_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 0.1 <= v['cvss_score'] < 4),
        "info_count": sum(1 for v in vulnerabilities.values() if not isinstance(v.get('cvss_score'), (int, float)) or v['cvss_score'] < 0.1),
        "kev_count": sum(1 for v in vulnerabilities.values() if v.get('cisa_kev') == 'Yes')
    }

    chart_b64 = generate_severity_chart(vulnerabilities)

    try:
        env = Environment(loader=FileSystemLoader(os.path.join(os.path.dirname(__file__), 'templates')))
        template = env.get_template('report_template.html')

        def get_score(v):
            s = v.get('cvss_score', 0)
            return float(s) if s != 'N/A' else 0

        sorted_vulns = dict(sorted(vulnerabilities.items(), key=lambda x: get_score(x[1]), reverse=True))

        html_out = template.render(
            vulnerabilities=sorted_vulns,
            cve_to_hosts=cve_to_hosts,
            stats=stats,
            chart_b64=chart_b64,
            date=datetime.now().strftime("%Y-%m-%d %H:%M"),
            scanner_type=scanner_type,
            target_file=os.path.basename(target_file)
        )
        HTML(string=html_out).write_pdf(output_file)
        print(f"[+] PDF intelligence report saved to: {output_file}")
    except Exception as e:
        print(f"[!] Error generating PDF report: {e}")

def save_html_report(vulnerabilities, cve_to_hosts, scanner_type, target_file, output_file):
    if not HAS_JINJA:
        print("[!] Jinja2 not installed. HTML report skipped.")
        return

    # Calculate statistics
    stats = {
        "total_hosts": len(set(h[0] for hosts in cve_to_hosts.values() for h in hosts)),
        "total_cves": len(vulnerabilities),
        "critical_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and v['cvss_score'] >= 9),
        "high_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 7 <= v['cvss_score'] < 9),
        "medium_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 4 <= v['cvss_score'] < 7),
        "low_count": sum(1 for v in vulnerabilities.values() if isinstance(v.get('cvss_score'), (int, float)) and 0.1 <= v['cvss_score'] < 4),
        "info_count": sum(1 for v in vulnerabilities.values() if not isinstance(v.get('cvss_score'), (int, float)) or v['cvss_score'] < 0.1),
        "kev_count": sum(1 for v in vulnerabilities.values() if v.get('cisa_kev') == 'Yes')
    }

    chart_b64 = generate_severity_chart(vulnerabilities)

    try:
        env = Environment(loader=FileSystemLoader(os.path.join(os.path.dirname(__file__), 'templates')))
        template = env.get_template('report_template.html')

        # Sort vulnerabilities by CVSS score descending
        def get_score(v):
            s = v.get('cvss_score', 0)
            return float(s) if s != 'N/A' else 0

        sorted_vulns = dict(sorted(vulnerabilities.items(), key=lambda x: get_score(x[1]), reverse=True))

        html_out = template.render(
            vulnerabilities=sorted_vulns,
            cve_to_hosts=cve_to_hosts,
            stats=stats,
            chart_b64=chart_b64,
            date=datetime.now().strftime("%Y-%m-%d %H:%M"),
            scanner_type=scanner_type,
            target_file=os.path.basename(target_file)
        )
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(html_out)
        print(f"[+] HTML intelligence report saved to: {output_file}")
    except Exception as e:
        print(f"[!] Error generating HTML report: {e}")

def get_severity(cvss_score) -> str:
    """Map CVSS score to severity level."""
    if cvss_score == "N/A": return "info"
    try:
        score = float(cvss_score)
        if score >= 9.0: return "critical"
        elif score >= 7.0: return "high"
        elif score >= 4.0: return "medium"
        elif score > 0.0: return "low"
        else: return "info"
    except (ValueError, TypeError):
        return "info"

def generate_report(scan_file: str, csv_file: str = None, github_token: str = None, nvd_key: str = None, vulners_key: str = None, circl_key: str = None, html_file: str = None, md_file: str = None, pdf_file: str = None, xlsx_file: str = None, docx_file: str = None, all_base: str = None, threads: int = 10, severity: str = None):
    # Handle the --all (-A) logic
    if all_base:
        csv_file = f"{all_base}.csv"
        html_file = f"{all_base}.html"
        md_file = f"{all_base}.md"
        pdf_file = f"{all_base}.pdf"
        xlsx_file = f"{all_base}.xlsx"
        docx_file = f"{all_base}.docx"
    
    # Severity filter processing
    allowed_severities = [s.strip().lower() for s in severity.split(',')] if severity else None

    # Smart Fallback: If absolutely no output is specified, default to CSV
    if not any([csv_file, html_file, md_file, pdf_file, xlsx_file, docx_file]):
        csv_file = "vulnerability_report.csv"
        print(f"[*] No output format specified. Defaulting to CSV: {csv_file}")

    scanner_type = scan2cve.detect_scanner(scan_file)
    if not scanner_type: print(f"[!] Could not detect scanner type for {scan_file}"); return
    print(f"[*] Detected scanner type: {scanner_type}")
    session = requests.Session()
    if scanner_type == 'nmap_xml': parser = scan2cve.NmapParser(scan_file, session)
    elif scanner_type == 'nessus_xml': parser = scan2cve.NessusXMLParser(scan_file)
    elif scanner_type == 'nessus_csv': parser = scan2cve.NessusCSVParser(scan_file)
    elif scanner_type == 'qualys_xml': parser = scan2cve.QualysXMLParser(scan_file)
    elif scanner_type == 'qualys_csv': parser = scan2cve.QualysCSVParser(scan_file)
    else: print("[!] Unsupported scanner type."); return

    parser.parse()
    if not parser.hosts: print("[!] No scan data found."); return
    
    # Resolve CPEs to CVEs for Nmap (concurrently)
    if scanner_type == 'nmap_xml':
        unique_valid_cpes = set()
        for _, host_data in parser.hosts.items():
            for s in host_data["services"]:
                for cpe in s.get('cpes', []):
                    if not scan2cve.is_broad_cpe(cpe):
                        unique_valid_cpes.add(cpe)
                    else:
                        print(f"  [-] Skipping broad/generic CPE: {cpe}")
        
        if unique_valid_cpes:
            print(f"[*] Fetching CVEs for {len(unique_valid_cpes)} unique specific CPEs using {threads} threads...")
            cpe_to_cves = {}
            with ThreadPoolExecutor(max_workers=threads) as executor:
                futures = {executor.submit(parser.get_cves_for_cpe, cpe, nvd_key): cpe for cpe in unique_valid_cpes}
                for f in as_completed(futures):
                    cpe = futures[f]
                    try:
                        cpe_to_cves[cpe] = f.result()
                    except Exception as e:
                        print(f"  [!] Error fetching {cpe}: {e}")
                        cpe_to_cves[cpe] = []
            
            # Map resolved CVEs back to services
            for _, host_data in parser.hosts.items():
                for s in host_data["services"]:
                    service_cves = []
                    for cpe in s.get('cpes', []):
                        if cpe in cpe_to_cves:
                            service_cves.extend(cpe_to_cves[cpe])
                    s['cves'] = set(service_cves)

    all_cve_ids = set()
    for _, host_data in parser.hosts.items():
        for s in host_data["services"]:
            all_cve_ids.update(s['cves'])

    cve_enrichment = get_detailed_cve_info(list(all_cve_ids), github_token, nvd_key, vulners_key, circl_key, max_workers=threads)
    
    # Apply Severity Filtering
    if allowed_severities:
        print(f"[*] Filtering results for severities: {', '.join(allowed_severities)}")
        filtered_cves = set()
        for cve_id, d in cve_enrichment.items():
            if get_severity(d.get('cvss_score', 'N/A')) in allowed_severities:
                filtered_cves.add(cve_id)
        
        # Prune services and enrichment
        for _, host_data in parser.hosts.items():
            for s in host_data["services"]:
                s['cves'] = {cid for cid in s['cves'] if cid in filtered_cves}
        
        cve_enrichment = {cid: data for cid, data in cve_enrichment.items() if cid in filtered_cves}
        all_cve_ids = filtered_cves

    headers = [
        "IP", "Hostname (FQDN)", "CVE ID", "Port", "Service",
        "Title", "Description", "CVSS Score", "CISA KEV", "EPSS Score",
        "Exploitability", "PoC Available", "PoC Link", "Threat Vector", "User Interaction", "Complexity", 
        "Affected Products", "Remediation", "References"
    ]

    # Prepare host mapping for reporting
    cve_to_hosts = {}
    for _, host_data in parser.hosts.items():
        ip, hostname = host_data.get('ip', ''), host_data.get('hostname', '')
        for s in host_data["services"]:
            port_proto = f"{s['port']}/{s['proto']}" if s['port'] != "0" else "Host-level"
            for cve_id in s['cves']:
                if cve_id not in cve_to_hosts: cve_to_hosts[cve_id] = []
                # Avoid duplicate (ip, hostname, port_proto, s['name']) for the same CVE
                host_entry = (ip, hostname, port_proto, s['name'])
                if host_entry not in cve_to_hosts[cve_id]:
                    cve_to_hosts[cve_id].append(host_entry)

    if csv_file:
        try:
            with open(csv_file, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                # CSV always defaults to host-based rows
                for _, host_data in sorted(parser.hosts.items()):
                    ip, hostname = host_data.get('ip', ''), host_data.get('hostname', '')
                    for s in host_data["services"]:
                        port_proto = f"{s['port']}/{s['proto']}" if s['port'] != "0" else "Host-level"
                        for cve_id in sorted(list(s['cves'])):
                            d = cve_enrichment.get(cve_id, {})
                            writer.writerow([
                                ip, hostname, cve_id, port_proto, s['name'],
                                d.get('title', 'N/A'), d.get('description', 'N/A'), d.get('cvss_score', 'N/A'),
                                d.get('cisa_kev', 'N/A'), d.get('epss_score', 'N/A'), d.get('exploitability', 'N/A'),
                                d.get('poc_available', 'N/A'), d.get('poc_link', 'N/A'), d.get('threat_vector', 'N/A'), d.get('user_interaction', 'N/A'),
                                d.get('attack_complexity', 'N/A'),
                                "; ".join([f"{a['product']} ({', '.join(a['versions'])})" for a in d.get('affected', [])]),
                                d.get('remediation', 'N/A'),
                                "; ".join(d.get('references', []))
                            ])
            print(f"[+] Detailed CSV report saved to: {csv_file}")
        except Exception as e: print(f"[!] Error writing CSV: {e}")

    if html_file:
        save_html_report(cve_enrichment, cve_to_hosts, scanner_type, scan_file, html_file)
    if md_file:
        save_markdown_report(cve_enrichment, cve_to_hosts, scanner_type, scan_file, md_file)
    if pdf_file:
        save_pdf_report(cve_enrichment, cve_to_hosts, scanner_type, scan_file, pdf_file)
    if xlsx_file:
        save_xlsx_report(cve_enrichment, cve_to_hosts, xlsx_file)
    if docx_file:
        save_docx_report(cve_enrichment, cve_to_hosts, scanner_type, scan_file, docx_file)

def main():
    p = argparse.ArgumentParser(description="VulnReport - Unified Scan Reporting Tool")
    p.add_argument("file", help="Input scan file (Nmap/Nessus/Qualys)")
    p.add_argument("-C", "--csv", help="Output CSV filename")
    p.add_argument("-H", "--html", help="Output HTML report filename")
    p.add_argument("-M", "--markdown", help="Output Markdown report filename")
    p.add_argument("-P", "--pdf", help="Output PDF report filename")
    p.add_argument("-X", "--xlsx", help="Output XLSX report filename")
    p.add_argument("-D", "--docx", help="Output Word DOCX report filename")
    p.add_argument("-A", "--all", help="Generate ALL formats using this base filename")
    p.add_argument("-t", "--threads", type=int, default=10, help="Number of concurrent threads (default: 10)")
    p.add_argument("-s", "--severity", help="Filter by severity: Info, Low, Medium, High, Critical (comma-separated)")
    args = p.parse_args()
    if not os.path.exists(args.file): print(f"[!] File not found: {args.file}"); return

    # Automatic key discovery
    github_token, nvd_key, vulners_key, circl_key = cve_lookup.get_keys()
    if nvd_key: print("[*] Using NVD API Key for accelerated lookups.")
    if github_token: print("[*] Using GitHub Token for PoC research.")
    if vulners_key: print("[*] Using Vulners API Key for enhanced intelligence.")
    if circl_key: print("[*] Using CIRCL API Key for accelerated lookups.")

    generate_report(args.file, args.csv, github_token, nvd_key, vulners_key, circl_key, args.html, args.markdown, args.pdf, args.xlsx, args.docx, args.all, args.threads, args.severity)

if __name__ == "__main__": main()
